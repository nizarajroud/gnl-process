"""Tests for the generic Doc→Meta tool (extraction + text splitting)."""
import pytest
from gnl_core.doc_extract import extract_text, SUPPORTED
from gnl_core.meta_export import build_parts_from_text, DEFAULT_GENERIC_PROMPT, DEFAULT_PROMPT


def test_extract_txt(tmp_path):
    p = tmp_path / "a.txt"; p.write_text("Hello world\nLine 2")
    assert "Hello world" in extract_text(str(p))


def test_extract_docx(tmp_path):
    from docx import Document
    d = Document(); d.add_paragraph("Para one"); d.add_paragraph("Para two")
    p = tmp_path / "a.docx"; d.save(str(p))
    t = extract_text(str(p))
    assert "Para one" in t and "Para two" in t


def test_unsupported_format(tmp_path):
    p = tmp_path / "a.exe"; p.write_bytes(b"MZ")
    with pytest.raises(ValueError):
        extract_text(str(p))


def test_empty_doc_raises(tmp_path):
    p = tmp_path / "a.txt"; p.write_text("   ")
    with pytest.raises(ValueError):
        extract_text(str(p))


def test_build_parts_from_text_single():
    parts = build_parts_from_text("Short paragraph.\n\nAnother one.", max_chars=65000)
    assert len(parts) == 1
    assert 'study/analysis partner' in parts[0]
    assert 'OF 1' in parts[0]


def test_build_parts_from_text_multiple():
    text = "\n\n".join(f"Paragraph {i} " + ("x" * 400) for i in range(50))
    parts = build_parts_from_text(text, max_chars=5000)
    assert len(parts) > 1
    for p in parts:
        assert len(p) <= 5000 + 1000
    assert 'study/analysis partner' in parts[0]
    assert all('study/analysis partner' not in p for p in parts[1:])


def test_generic_prompt_distinct_from_exam():
    # Guarantee the generic tool does NOT reuse the exam prompt.
    assert DEFAULT_GENERIC_PROMPT != DEFAULT_PROMPT
    assert 'senior AWS' not in DEFAULT_GENERIC_PROMPT


def test_oversized_paragraph_is_split():
    big = "word " * 30000  # one giant paragraph
    parts = build_parts_from_text(big, max_chars=10000)
    assert len(parts) > 1
    for p in parts:
        assert len(p) <= 10000 + 1000
