"""Tests for the plug-and-play exam format registry."""
import os
from docx import Document
import pytest

from gnl_core.exam_formats import (
    detect_format, get_format_by_name, FORMATS,
    AwsPracticeExamFormat, DojoFormat,
)
from gnl_core.exam_formats.base import normalize_question
from gnl_core.exam_formats.pipeline import pivot_to_answers, pivot_to_markdown


def _make_practice_docx(path):
    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("Multiple Choice")
    doc.add_paragraph("Answer status:")
    doc.add_paragraph("Incorrect")
    doc.add_paragraph("Question")
    doc.add_paragraph("Which storage is cheapest for archival?")
    doc.add_paragraph("Answer options")
    t = doc.add_table(rows=3, cols=4)
    h = t.rows[0].cells
    h[0].text, h[1].text, h[2].text, h[3].text = 'Option', 'Correct answer', 'Your selection', 'Rationale'
    r1 = t.rows[1].cells
    r1[0].text, r1[1].text, r1[2].text, r1[3].text = 'A. S3 Standard', '', 'Selected', 'Not the cheapest.'
    r2 = t.rows[2].cells
    r2[0].text, r2[1].text, r2[2].text, r2[3].text = 'B. Glacier Deep Archive', 'Correct', 'Not selected', 'Cheapest for archival.'
    doc.save(str(path))


def _make_dojo_docx(path):
    doc = Document()
    doc.add_paragraph("1. Question")
    doc.add_paragraph("A company needs to store data. Which option?")
    doc.add_paragraph("Use S3")
    doc.add_paragraph("Use EBS")
    doc.add_paragraph("Correct options: Use S3")
    doc.add_paragraph("References: https://aws.amazon.com/s3")
    doc.save(str(path))


def test_detect_practice_exam(tmp_path):
    p = tmp_path / "prac.docx"; _make_practice_docx(p)
    fmt = detect_format(str(p))
    assert fmt is not None and fmt.name == 'aws_practice_exam'


def test_detect_dojo(tmp_path):
    p = tmp_path / "dojo-something.docx"; _make_dojo_docx(p)
    fmt = detect_format(str(p))
    assert fmt is not None and fmt.name == 'dojo'


def test_practice_exam_pivot(tmp_path):
    p = tmp_path / "prac.docx"; _make_practice_docx(p)
    qs = AwsPracticeExamFormat().parse(str(p))
    assert len(qs) == 1
    q = qs[0]
    assert q['num'] == 1 and q['status'] == 'Incorrect'
    a, b = q['options']
    assert a['letter'] == 'A' and a['selected'] and not a['correct']
    assert b['letter'] == 'B' and b['correct'] and not b['selected']
    assert 'cheapest' in b['rationale'].lower()


def test_pivot_to_answers_and_markdown():
    pivot = [normalize_question(1, "Stmt?", [
        {'letter': 'A', 'text': 'Opt A', 'correct': False},
        {'letter': 'B', 'text': 'Opt B', 'correct': True, 'rationale': 'because B'},
    ])]
    ans = pivot_to_answers(pivot)
    assert ans['1']['correct'] == ['Opt B']
    assert ans['1']['type'] == 'single'
    md = pivot_to_markdown(pivot, "MyExam")
    assert '## Question 1:' in md and '- Opt A' in md and 'Explanations:' in md


def test_multi_answer_type():
    pivot = [normalize_question(1, "S?", [
        {'letter': 'A', 'text': 'A', 'correct': True},
        {'letter': 'B', 'text': 'B', 'correct': True},
        {'letter': 'C', 'text': 'C', 'correct': False},
    ])]
    assert pivot_to_answers(pivot)['1']['type'] == 'multi'


def test_registry_order_practice_before_dojo():
    # A practice-exam doc must never be captured by dojo (tables win).
    assert FORMATS[0].name == 'aws_practice_exam'


def test_non_docx_detected_as_none(tmp_path):
    p = tmp_path / "notes.txt"; p.write_text("hello")
    assert detect_format(str(p)) is None
