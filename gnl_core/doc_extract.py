"""
Generic document text extraction — multi-format.

Extracts plain text from an uploaded document regardless of format, so it can
be fed to the generic Meta export (split into paste-ready parts).

Supported: .pdf, .docx, .pptx, .xlsx, .txt, .md (and .markdown).
Unknown/binary formats raise a clear ValueError.

This module is INDEPENDENT of the exam pipeline — it does not import or modify
anything exam-related.
"""

import os
from pathlib import Path

SUPPORTED = ('.pdf', '.docx', '.pptx', '.xlsx', '.txt', '.md', '.markdown')


def _extract_pdf(path):
    try:
        from pypdf import PdfReader
    except Exception:
        from PyPDF2 import PdfReader
    reader = PdfReader(path)
    out = []
    for i, page in enumerate(reader.pages, start=1):
        txt = page.extract_text() or ''
        if txt.strip():
            out.append(txt.strip())
    return "\n\n".join(out)


def _extract_docx(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # include tables
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines = [f"## Slide {idx}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(t)
        if len(lines) > 1:
            out.append("\n".join(lines))
    return "\n\n".join(out)


def _extract_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _extract_text(path):
    return Path(path).read_text(encoding='utf-8', errors='ignore')


_EXTRACTORS = {
    '.pdf': _extract_pdf,
    '.docx': _extract_docx,
    '.pptx': _extract_pptx,
    '.xlsx': _extract_xlsx,
    '.txt': _extract_text,
    '.md': _extract_text,
    '.markdown': _extract_text,
}


def extract_text(path):
    """Extract plain text from a document. Raises ValueError for unsupported
    formats, or on extraction failure with a clear message."""
    ext = Path(path).suffix.lower()
    fn = _EXTRACTORS.get(ext)
    if fn is None:
        raise ValueError(
            f"Format non supporté: {ext or '(aucune extension)'}. "
            f"Formats acceptés: {', '.join(SUPPORTED)}")
    try:
        text = fn(path)
    except Exception as e:
        raise ValueError(f"Extraction échouée ({ext}): {str(e)[:120]}")
    text = (text or '').strip()
    if not text:
        raise ValueError(f"Aucun texte extrait de ce {ext} (document vide ou scanné/image ?)")
    return text
