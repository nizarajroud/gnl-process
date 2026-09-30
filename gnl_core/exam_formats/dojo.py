"""
DOJO Timed Mode format.

Answers are NOT explicit in the document, so this format resolves correct
answers via the existing 3-tier highlight (NotebookLM -> Bedrock -> regex).

Zero-regression guarantee: this parser REUSES the existing, unchanged pipeline
functions (step1_format, step2b_full_markdown, step3_highlight,
parse_exam_questions). It only ADAPTS their output into the normalized pivot —
it does not reimplement any DOJO logic.
"""

import re
from .base import ExamFormat, normalize_question


class DojoFormat(ExamFormat):
    name = 'dojo'
    label = 'DOJO Timed Mode'

    def detect(self, path: str) -> bool:
        """DOJO docs are paragraph-based (no option tables) and contain
        'Question N' headers + a 'References:' or DOJO-style layout.
        We also match on filename as a strong hint."""
        if not str(path).lower().endswith('.docx'):
            return False
        try:
            from docx import Document
            doc = Document(path)
            # If it has practice-exam option tables, it's NOT dojo.
            for tbl in doc.tables:
                if tbl.rows:
                    header = [c.text.strip().lower() for c in tbl.rows[0].cells]
                    if 'option' in header and 'correct answer' in header:
                        return False
            text = "\n".join(p.text for p in doc.paragraphs[:400])
            # DOJO uses either "Question N" (post-format) or "N. Question"
            # (raw origin export).
            has_questions = bool(re.search(r'Question\s+\d+', text)) \
                or bool(re.search(r'^\s*\d+\.\s*Question\s*$', text, re.M))
            # DOJO hallmark: "References:" blocks and/or "Overall explanation",
            # or the filename explicitly mentions dojo.
            dojo_hint = ('dojo' in str(path).lower()
                         or 'References:' in text
                         or 'Overall explanation' in text)
            return has_questions and dojo_hint
        except Exception:
            return False

    def parse(self, path: str, on_progress=None, *, theme=None, subtheme=None) -> list:
        """Run the existing DOJO pipeline (unchanged) and adapt to the pivot.

        Needs theme/subtheme for path resolution (the existing steps write into
        the exam asset tree). Returns the normalized pivot list.
        """
        from gnl_core.exams import (
            step1_format, step2b_full_markdown, step3_highlight,
        )
        from gnl_core.anki_review import parse_exam_questions

        origin = 'dojo'
        # Step 1 & 2: origin -> word -> markdown (unchanged existing funcs)
        word_path = step1_format(path, theme, subtheme, origin, on_progress=on_progress)
        md_path = step2b_full_markdown(word_path, theme, subtheme, on_progress=on_progress)
        # Step 3: resolve correct answers (NLM/Bedrock/regex) — unchanged
        answers = step3_highlight(md_path, on_progress=on_progress)
        # Question texts + explanations from the generated markdown
        blocks = parse_exam_questions(md_path)

        # Re-read option texts from the markdown (same source step5_anki uses)
        import re as _re
        from pathlib import Path
        md = Path(md_path).read_text(encoding='utf-8', errors='ignore')
        parts = _re.split(r'## Question\s+\d+:', md)
        q_headers = _re.findall(r'## (Question\s+(\d+):)', md)
        opts_by_num = {}
        for idx, (_, num) in enumerate(q_headers):
            if idx + 1 < len(parts):
                block = parts[idx + 1]
                opts = [ln.strip()[2:].strip()
                        for ln in block.split('\n') if ln.strip().startswith('- ')]
                opts_by_num[num] = opts

        questions = []
        for num in sorted(answers.keys(), key=lambda x: int(x)):
            entry = answers[num]
            correct = entry.get('correct', [])
            correct_norm = [c.lower().strip() for c in correct]
            opts = opts_by_num.get(str(num), entry.get('options', []))
            options = []
            for opt in opts:
                on = opt.lower().strip()
                is_c = any(on == cn or (len(cn) > 20 and cn in on)
                           or (len(on) > 20 and on in cn) for cn in correct_norm)
                options.append({'letter': '', 'text': opt, 'correct': is_c,
                                'selected': False, 'rationale': ''})
            expl = blocks.get(int(num), {}).get('explanation', '') if blocks else ''
            body = blocks.get(int(num), {}).get('body', '') if blocks else ''
            questions.append(normalize_question(
                num, body, options, explanation=expl))
        if on_progress:
            on_progress(f"DOJO: {len(questions)} questions")
        # Stash the md_path so the pipeline can reuse it (avoids recompute)
        self._last_md_path = md_path
        self._last_answers = answers
        return questions
