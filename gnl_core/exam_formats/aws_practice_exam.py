"""
AWS Practice Exam format.

Structure (per question, e.g. sap-aws-practice-exam.docx):
    Paragraphs: "Question N" / "Multiple Choice" / "Answer status:" / <status>
                / "Question" / <statement paras...> / "Answer options"
    Table (4 cols): Option | Correct answer | Your selection | Rationale

Answers are EXPLICIT in the document (no Bedrock needed):
    - Correct answer column == "Correct"       -> correct option
    - Your selection column == "Selected"      -> the user's own answer
    - Rationale column                         -> per-option explanation
"""

import re
from .base import ExamFormat, normalize_question


class AwsPracticeExamFormat(ExamFormat):
    name = 'aws_practice_exam'
    label = 'AWS Practice Exam'

    def _iter_blocks(self, doc):
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        for child in doc.element.body.iterchildren():
            if child.tag == qn('w:p'):
                yield ('p', Paragraph(child, doc))
            elif child.tag == qn('w:tbl'):
                yield ('tbl', Table(child, doc))

    def detect(self, path: str) -> bool:
        """True if the doc has option tables with the practice-exam header."""
        if not str(path).lower().endswith('.docx'):
            return False
        try:
            from docx import Document
            doc = Document(path)
            for tbl in doc.tables:
                if not tbl.rows:
                    continue
                header = [c.text.strip().lower() for c in tbl.rows[0].cells]
                if 'option' in header and 'correct answer' in header \
                        and 'your selection' in header:
                    return True
            return False
        except Exception:
            return False

    def parse(self, path: str, on_progress=None) -> list:
        from docx import Document
        doc = Document(path)
        blocks = list(self._iter_blocks(doc))
        questions = []
        i, n = 0, len(blocks)
        while i < n:
            kind, el = blocks[i]
            if kind == 'p' and re.match(r'^Question\s+\d+\s*$', el.text.strip()):
                num = int(re.search(r'\d+', el.text).group())
                status = ''
                statement_lines = []
                in_statement = False
                j = i + 1
                while j < n:
                    k2, e2 = blocks[j]
                    if k2 == 'tbl':
                        break
                    if k2 == 'p':
                        t = e2.text.strip()
                        if re.match(r'^Question\s+\d+\s*$', t):
                            break
                        if t in ('Correct', 'Incorrect') and status == '':
                            status = t
                        elif t == 'Question':
                            in_statement = True
                        elif t == 'Answer options':
                            in_statement = False
                        elif in_statement and t:
                            statement_lines.append(t)
                    j += 1
                options = []
                if j < n and blocks[j][0] == 'tbl':
                    tbl = blocks[j][1]
                    rows = tbl.rows
                    if rows and rows[0].cells[0].text.strip().lower() == 'option':
                        for r in rows[1:]:
                            cells = r.cells
                            opt_text = cells[0].text.strip()
                            m = re.match(r'^([A-Z])\.\s*(.*)', opt_text, re.S)
                            letter = m.group(1) if m else ''
                            text = (m.group(2).strip() if m else opt_text)
                            options.append({
                                'letter': letter,
                                'text': text,
                                'correct': cells[1].text.strip().lower() == 'correct',
                                'selected': cells[2].text.strip().lower() == 'selected',
                                'rationale': cells[3].text.strip() if len(cells) > 3 else '',
                            })
                    i = j
                questions.append(normalize_question(
                    num, '\n'.join(statement_lines), options, status=status))
            i += 1
        if on_progress:
            on_progress(f"AWS Practice Exam: {len(questions)} questions")
        return questions
